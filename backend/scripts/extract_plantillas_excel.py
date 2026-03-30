# -*- coding: utf-8 -*-
"""Extrae plantillas de entregables desde el Excel de matrices para generar SQL."""
import argparse
import re
from pathlib import Path

try:
    import openpyxl
except ImportError:
    raise SystemExit("Instale openpyxl: pip install openpyxl")

UBIC_RE = re.compile(
    r"(Febrero|Marzo|Abril|Mayo)\s*-\s*Semana\s*(\d+)",
    re.IGNORECASE,
)
MES_MAP = {
    "febrero": "febrero",
    "marzo": "marzo",
    "abril": "abril",
    "mayo": "mayo",
}


def sql_val(v):
    if v is None:
        return "NULL"
    if isinstance(v, bool):
        return "TRUE" if v else "FALSE"
    if isinstance(v, (int, float)):
        return str(v)
    s = str(v).replace("'", "''").strip()
    return f"'{s}'"


def sql_escape_text(v):
    if v is None:
        return ""
    return str(v).replace("'", "''")


def norm_cell(val):
    if val is None:
        return None
    if isinstance(val, str):
        return val.strip()
    if isinstance(val, (int, float)):
        return val
    return val


def is_section_header(b_str: str) -> bool:
    if not b_str or len(b_str) < 3:
        return False
    if re.match(r"^\d+\.\s", b_str.strip()):
        return True
    u = b_str.upper()
    if "UBICACI" in u and "MS" in u:
        return False
    if "LISTADO" in u:
        return False
    if "CAPACIT" in u or "CONVENIO" in u or "ACTIVIDAD" in u or "PROYECTO" in u:
        if len(b_str) > 15 and "SEMANA" not in b_str[:30]:
            return True
    return False


def find_sheet(wb, sheet_contains, sheet_exact):
    if sheet_exact:
        for name in wb.sheetnames:
            if name.strip() == sheet_exact.strip():
                return wb[name]
        raise SystemExit(f"Hoja exacta no encontrada: {sheet_exact!r}")
    if sheet_contains:
        c = sheet_contains.strip().lower()
        for name in wb.sheetnames:
            if c in name.lower():
                return wb[name]
        raise SystemExit(f"No hay hoja que contenga {sheet_contains!r}")
    raise SystemExit("Especifique --sheet-contains o --sheet-exact")


def extract_rows(sh, grupo_id: int):
    last_ubic = None
    last_fase = None
    rows_out = []

    for r in range(1, (sh.max_row or 0) + 1):
        b = norm_cell(sh.cell(r, 2).value)
        c = norm_cell(sh.cell(r, 3).value)
        d = norm_cell(sh.cell(r, 4).value)
        e = norm_cell(sh.cell(r, 5).value)
        f = norm_cell(sh.cell(r, 6).value)
        h = norm_cell(sh.cell(r, 8).value)

        b_str = ""
        if isinstance(b, str):
            b_str = b.strip()
        elif b is not None:
            b_str = str(b).strip()

        if isinstance(b, str) and b_str:
            if is_section_header(b_str):
                last_fase = b_str
            mubic = UBIC_RE.search(b_str.replace("\n", " "))
            if mubic:
                first_line = b_str.split("\n")[0].strip()
                if UBIC_RE.search(first_line):
                    last_ubic = first_line

        if isinstance(c, float) and c == int(c):
            c = int(c)
        if not isinstance(c, int):
            continue
        if d is None or (isinstance(d, str) and not d.strip()):
            continue
        if isinstance(d, str) and d.strip().lower() == "entregable":
            continue

        ubic_ref = last_ubic
        if isinstance(b, str) and b_str:
            m = UBIC_RE.search(b_str.replace("\n", " "))
            if m:
                ubic_ref = b_str.split("\n")[0].strip()

        if not ubic_ref:
            continue

        m2 = UBIC_RE.search(ubic_ref.replace("\n", " "))
        if not m2:
            continue

        mes_name = m2.group(1).lower()
        mes_db = MES_MAP.get(mes_name)
        if not mes_db:
            continue
        semana_numero = int(m2.group(2))

        entregable = d.strip() if isinstance(d, str) else str(d)
        desc_ev = (
            (e.strip() if isinstance(e, str) else str(e)) if e is not None else ""
        )

        if f is None:
            horas = 0.0
        elif isinstance(f, (int, float)):
            horas = float(f)
        else:
            try:
                horas = float(str(f).replace(",", "."))
            except ValueError:
                horas = 0.0

        if isinstance(h, str) and h.strip():
            fase_val = h.strip()
        else:
            fase_val = last_fase or ""

        categoria_db = ubic_ref
        dias_ini = (semana_numero - 1) * 7
        dias_fin = semana_numero * 7 - 1

        rows_out.append(
            {
                "numero": c,
                "categoria": categoria_db,
                "fase": fase_val,
                "mes": mes_db,
                "semana_numero": semana_numero,
                "entregable": entregable,
                "descripcion_evidencia": desc_ev,
                "horas": horas,
                "dias_inicio_desde_feb": dias_ini,
                "dias_fin_desde_feb": dias_fin,
            }
        )

    return rows_out


def emit_sql(rows, grupo_id: int, fh, delete_first: bool):
    if delete_first:
        fh.write(f"DELETE FROM plantilla_entregables WHERE grupo_id = {grupo_id};\n\n")
    if not rows:
        fh.write("-- Sin filas extraídas\n")
        return
    fh.write(
        "INSERT INTO plantilla_entregables "
        "(grupo_id, numero, categoria, fase, mes, semana_numero, entregable, descripcion_evidencia, horas, dias_inicio_desde_feb, dias_fin_desde_feb) VALUES\n"
    )
    parts = []
    for row in rows:
        parts.append(
            "({gid}, {num}, {cat}, {fase}, {mes}, {sem}, {ent}, {desc}, {hrs}, {di}, {df})".format(
                gid=grupo_id,
                num=row["numero"],
                cat=sql_val(row["categoria"]),
                fase=sql_val(row["fase"]),
                mes=sql_val(row["mes"]),
                sem=row["semana_numero"],
                ent=sql_val(row["entregable"]),
                desc=sql_val(row["descripcion_evidencia"]),
                hrs=row["horas"],
                di=row["dias_inicio_desde_feb"],
                df=row["dias_fin_desde_feb"],
            )
        )
    fh.write(",\n".join(parts))
    fh.write(";\n")


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--excel", required=True, help="Ruta al .xlsx")
    ap.add_argument("--grupo-id", type=int, required=True)
    ap.add_argument("--output", required=True, help="Archivo .sql de salida")
    ap.add_argument("--sheet-contains", default=None)
    ap.add_argument("--sheet-exact", default=None)
    ap.add_argument("--delete-first", action="store_true")
    args = ap.parse_args()

    path = Path(args.excel)
    if not path.is_file():
        raise SystemExit(f"No existe: {path}")

    wb = openpyxl.load_workbook(str(path), data_only=True)
    sh = find_sheet(wb, args.sheet_contains, args.sheet_exact)
    rows = extract_rows(sh, args.grupo_id)

    out = Path(args.output)
    out.parent.mkdir(parents=True, exist_ok=True)
    with open(out, "w", encoding="utf-8") as fh:
        fh.write(f"-- Generado por extract_plantillas_excel.py ({len(rows)} filas)\n")
        emit_sql(rows, args.grupo_id, fh, args.delete_first)

    print(f"OK: {len(rows)} filas -> {out}")


if __name__ == "__main__":
    main()
